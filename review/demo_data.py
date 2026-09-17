"""Bundled sample data for `ta review --demo` and the TUI preview renderer.

Everything here is fake: no network access, no credentials, no submissions
from real students.
"""
from __future__ import annotations

import tempfile
from pathlib import Path

from models import CriterionScore, ScoringResult, UncertainPart
from review.spreadsheet import export

DEMO_RUBRIC = """# Homework 1 Rubric (100 points)

## Problem 1.2 Selection Sort (40 pts)
- Correct algorithm (20 pts)
- Correct comparison count (10 pts)
- Correct complexity analysis (10 pts)

## Problem 1.6 Matrix Multiplication (40 pts)
- Correct algorithm (20 pts)
- Correct multiplication count (10 pts)
- Correct addition count (10 pts)

## Problem 1.9 Recurrence Solving (20 pts)
- Correct master theorem application (10 pts)
- Correct final bound (10 pts)
"""

_DEMO_FILES = [
    ("2023001001", "Alice_Zhang", "Problem 1.2: O(n^2) comparisons. Problem 1.6: n^3 multiplications.\n"),
    ("2023001002", "Bob_Li", "Full solutions for problems 1.2, 1.6 and 1.9.\n"),
    ("2023001003", "Carol_Wang", "Handwritten solution, scanned. See attachment pages.\n"),
    ("2023001004", "David_Chen", "Problem 1.9 solved with the substitution method instead.\n"),
]


def demo_results() -> list[ScoringResult]:
    """Four fake scoring results covering the interesting review states."""
    return [
        ScoringResult(
            student_id="2023001001",
            student_name="Alice Zhang",
            assignment_id="423829",
            total_score=95.0,
            total_max=100.0,
            confidence=0.93,
            breakdown=[
                CriterionScore(criterion="1.2 Selection Sort",
                               points_awarded=40.0, points_max=40.0,
                               reasoning="Algorithm and comparison count are correct."),
                CriterionScore(criterion="1.6 Matrix Multiplication",
                               points_awarded=40.0, points_max=40.0,
                               reasoning="All counts correct."),
                CriterionScore(criterion="1.9 Recurrence Solving",
                               points_awarded=15.0, points_max=20.0,
                               reasoning="deduct 5 pts: the final bound is stated without justification."),
            ],
            llm_reasoning="Strong submission. One deduction: missing justification for the final bound.",
        ),
        ScoringResult(
            student_id="2023001002",
            student_name="Bob Li",
            assignment_id="423829",
            total_score=100.0,
            total_max=100.0,
            confidence=0.97,
            breakdown=[
                CriterionScore(criterion="1.2 Selection Sort",
                               points_awarded=40.0, points_max=40.0,
                               reasoning="Correct in every respect."),
                CriterionScore(criterion="1.6 Matrix Multiplication",
                               points_awarded=40.0, points_max=40.0,
                               reasoning="Correct in every respect."),
                CriterionScore(criterion="1.9 Recurrence Solving",
                               points_awarded=20.0, points_max=20.0,
                               reasoning="Correct master theorem application and bound."),
            ],
            llm_reasoning="Perfect submission.",
        ),
        ScoringResult(
            student_id="2023001003",
            student_name="Carol Wang",
            assignment_id="423829",
            total_score=78.0,
            total_max=100.0,
            confidence=0.58,
            needs_review=True,
            breakdown=[
                CriterionScore(criterion="1.2 Selection Sort",
                               points_awarded=34.0, points_max=40.0,
                               reasoning="deduct 6 pts: comparison count off by one."),
                CriterionScore(criterion="1.6 Matrix Multiplication",
                               points_awarded=30.0, points_max=40.0,
                               reasoning="deduct 10 pts: addition count is hard to read."),
                CriterionScore(criterion="1.9 Recurrence Solving",
                               points_awarded=14.0, points_max=20.0,
                               reasoning="deduct 6 pts: case 2 bound missing."),
            ],
            uncertain_parts=[
                UncertainPart(description="Handwriting on the recurrence derivation is hard to read",
                              suggested_score=18.0, suggested_max=20.0),
            ],
            llm_reasoning="Scanned handwriting makes parts of the solution ambiguous. Flagged for human review.",
        ),
        ScoringResult(
            student_id="2023001004",
            student_name="David Chen",
            assignment_id="423829",
            total_score=84.0,
            total_max=100.0,
            confidence=0.88,
            breakdown=[
                CriterionScore(criterion="1.2 Selection Sort",
                               points_awarded=38.0, points_max=40.0,
                               reasoning="deduct 2 pts: complexity analysis omits the best case."),
                CriterionScore(criterion="1.6 Matrix Multiplication",
                               points_awarded=40.0, points_max=40.0,
                               reasoning="Correct."),
                CriterionScore(criterion="1.9 Recurrence Solving",
                               points_awarded=6.0, points_max=20.0,
                               reasoning="deduct 14 pts: substitution method does not prove the tight bound."),
            ],
            llm_reasoning="Mostly correct; the recurrence solution uses the wrong method for the required bound.",
        ),
    ]


def create_demo_workspace(root: Path | None = None) -> tuple[Path, Path, Path]:
    """Create a demo workspace and return (scores.xlsx, submissions/, rubric.md).

    By default the files are written to a fresh temporary directory so nothing
    in the repository is touched.
    """
    if root is None:
        root = Path(tempfile.mkdtemp(prefix="pku-ta-demo-"))
    else:
        root = Path(root)
        root.mkdir(parents=True, exist_ok=True)

    scores = root / "scores.xlsx"
    export(demo_results(), scores)

    # Pre-approve the perfect submission so the approved state is visible too.
    import openpyxl

    wb = openpyxl.load_workbook(scores)
    ws = wb.active
    headers = {cell.value: i for i, cell in enumerate(ws[1])}
    for row in ws.iter_rows(min_row=2):
        if row[headers["student_id"]].value == "2023001002":
            row[headers["approved"]].value = "YES"
    wb.save(scores)

    homework_dir = root / "submissions" / "Homework_1"
    homework_dir.mkdir(parents=True, exist_ok=True)
    for student_id, name, text in _DEMO_FILES:
        (homework_dir / f"{student_id}_{name}.txt").write_text(text, encoding="utf-8")

    rubric = root / "rubric.md"
    rubric.write_text(DEMO_RUBRIC, encoding="utf-8")

    return scores, root / "submissions", rubric
